from pathlib import Path
from zipfile import ZipFile

import pandas as pd
import pytest
from openpyxl import load_workbook

from okpop import area, report

REPORT_DB_PATH = Path(
    "tests/fixtures/report/report.sqlite"
)

TEMPLATE_PATH = Path(
    "assets/templates/populaiton_report_template_0001.xlsx"
)


# --------------------------------------------------
# read_report_data
# --------------------------------------------------


def test_read_report_data():
    """人口と年齢区分がそろう最新月からデータを取得できる。"""

    df = report.read_report_data(
        area_name="大阪市福島区",
        end_date="2026-10-01",
        db_path=REPORT_DB_PATH,
    )

    assert not df.empty

    # 2026-09は年齢区分がないため、
    # 2026-08が最新月になる。
    assert (
        df["observation_date"].iloc[0]
        == pd.Timestamp("2026-08-01")
    )

    assert pd.notna(
        df["total_population"].iloc[0]
    )

    assert pd.notna(
        df["total_age_00_04"].iloc[0]
    )


# --------------------------------------------------
# validate_report_data
# --------------------------------------------------


def test_validate_report_data():
    """正常な実データなら検証を通過する。"""

    df = report.read_report_data(
        area_name="大阪市福島区",
        end_date="2026-10-01",
        db_path=REPORT_DB_PATH,
    )

    report.validate_report_data(df)


def test_validate_report_data_wrong_column_count():
    """76列でなければエラー。"""

    df = report.read_report_data(
        area_name="大阪市福島区",
        end_date="2026-10-01",
        db_path=REPORT_DB_PATH,
    )

    # 1列削除して75列にする
    df = df.drop(
        columns="total_population"
    )

    assert df.shape[1] == 75

    with pytest.raises(
        ValueError,
        match="76列",
    ):
        report.validate_report_data(df)


def test_validate_report_data_not_enough():
    """73か月未満ならエラー。"""

    df = report.read_report_data(
        area_name="大阪市福島区",
        end_date="2026-10-01",
        db_path=REPORT_DB_PATH,
    )

    # 最新72か月だけにする
    df = df.head(72)

    assert len(df) == 72

    with pytest.raises(
        ValueError,
        match="73か月",
    ):
        report.validate_report_data(df)


def test_validate_report_data_missing_month():
    """73件あっても途中の月が抜けていればエラー。"""

    df = report.read_report_data(
        area_name="大阪市福島区",
        end_date="2026-10-01",
        db_path=REPORT_DB_PATH,
    )

    # 最新73か月だけを使用する
    df = df.head(73).copy()

    # 途中の1か月を削除する
    df = df.drop(
        index=df.index[10]
    )

    # 代わりに73か月より前の1行を追加して、
    # 行数自体は73行に戻す。
    original_df = report.read_report_data(
        area_name="大阪市福島区",
        end_date="2026-10-01",
        db_path=REPORT_DB_PATH,
    )

    df = pd.concat(
        [
            df,
            original_df.iloc[[73]],
        ],
        ignore_index=True,
    )

    assert len(df) == 73

    with pytest.raises(
        ValueError,
        match="連続していません",
    ):
        report.validate_report_data(df)


def test_validate_report_data_missing_value():
    """許容されていない欠損値があればエラー。"""

    df = report.read_report_data(
        area_name="大阪市福島区",
        end_date="2026-10-01",
        db_path=REPORT_DB_PATH,
    )

    # 本来存在する人口を意図的に欠損させる
    df.loc[
        df.index[0],
        "total_population",
    ] = None

    with pytest.raises(
        ValueError,
        match="欠損値",
    ):
        report.validate_report_data(df)


# --------------------------------------------------
# create_report
# --------------------------------------------------


def test_create_report(tmp_path):
    """実DBからExcelレポートを作成できる。"""

    output_path = tmp_path / "report.xlsx"

    report.create_report(
        area_name="大阪市福島区",
        end_date="2026-10-01",
        output_path=output_path,
        template_path=TEMPLATE_PATH,
        db_path=REPORT_DB_PATH,
    )

    assert output_path.exists()

    wb = load_workbook(
        output_path,
        data_only=False,
    )

    assert "report" in wb.sheetnames
    assert "source" in wb.sheetnames

    ws_source = wb["source"]

    # 2026-09は年齢区分がないため、
    # sourceの最新月は2026-08になる。
    assert (
        ws_source["A2"].value.date()
        == pd.Timestamp("2026-08-01").date()
    )


# --------------------------------------------------
# create_report_archive
# --------------------------------------------------


def test_create_report_archive(
    tmp_path,
    monkeypatch,
):
    """全地域のレポートをZIP化できることを確認する。"""

    # create_report() をテスト用の偽物にする
    def fake_create_report(
        area_name,
        end_date,
        output_path,
        template_path,
        db_path,
    ):
        Path(output_path).write_text(
            area_name,
            encoding="utf-8",
        )
    
    monkeypatch.setattr(
        report,
        "get_latest_archive_date",
        lambda db_path: pd.Timestamp(
            "2025-09-01"
        ),
    )

    monkeypatch.setattr(
        report,
        "create_report",
        fake_create_report,
    )

    archive_file = report.create_report_archive(
        end_date="2025-08-01",
        output_dir=tmp_path,
        template_path="dummy.xlsx",
        db_path="dummy.sqlite",
    )

    # ZIPが作られた
    assert archive_file.exists()
    assert archive_file.suffix == ".zip"

    # ZIP化前のディレクトリは削除された
    report_dir = archive_file.with_suffix("")

    assert not report_dir.exists()

    # ZIPの中身を確認
    with ZipFile(archive_file) as zip_file:
        filenames = zip_file.namelist()

    # 全地域分存在する
    assert len(filenames) == len(
        area.AREA_NAME_V002_COL_NAME
    )

    # すべてxlsx
    assert all(
        filename.endswith(".xlsx")
        for filename in filenames
    )

    # 福島区も含まれる
    assert any(
        filename.startswith("大阪市福島区_")
        for filename in filenames
    )


def test_create_report_archive_reuses_existing_zip(
    monkeypatch,
    tmp_path,
):
    archive_file = (
        tmp_path / "okpop_report_202509.zip"
    )

    archive_file.write_bytes(
        b"existing zip"
    )

    monkeypatch.setattr(
        report,
        "get_latest_archive_date",
        lambda db_path: pd.Timestamp(
            "2025-09-01"
        ),
    )

    def fail_create_report(*args, **kwargs):
        raise AssertionError(
            "create_report should not be called"
        )

    monkeypatch.setattr(
        report,
        "create_report",
        fail_create_report,
    )

    result = report.create_report_archive(
        output_dir=tmp_path,
        template_path=tmp_path / "template.xlsx",
        db_path=tmp_path / "test.sqlite",
    )

    assert result == archive_file

    assert archive_file.read_bytes() == (
        b"existing zip"
    )
